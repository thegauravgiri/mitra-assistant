import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(0.92)
    section.bottom_margin = Inches(0.92)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("ITNET312A Major Individual Project | Project Plan")
    hrun.font.name = "Calibri"
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(100, 116, 139)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Sneha Neupane (900008705) • TAFE NSW Higher Education")
    frun.font.name = "Calibri"
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = RGBColor(100, 116, 139)

C_NAVY = RGBColor(30, 58, 138)
C_TEAL = RGBColor(15, 118, 110)
C_CHARCOAL = RGBColor(31, 41, 55)
C_BODY = RGBColor(51, 65, 85)
C_MUTED = RGBColor(100, 116, 139)

normal_style = doc.styles['Normal']
normal_style.font.name = 'Calibri'
normal_style.font.size = Pt(11)
normal_style.font.color.rgb = C_BODY
normal_style.paragraph_format.line_spacing = 1.08
normal_style.paragraph_format.space_after = Pt(2.0)

def add_title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.0)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(19)
    run.font.bold = True
    run.font.color.rgb = C_NAVY
    return p

def add_subtitle(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.color.rgb = C_TEAL
    run.font.italic = True
    return p

def add_meta_box(meta_dict):
    table = doc.add_table(rows=len(meta_dict), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    col_widths = [Inches(1.8), Inches(4.7)]
    for i, (k, v) in enumerate(meta_dict.items()):
        row = table.rows[i]
        cell0 = row.cells[0]
        cell0.width = col_widths[0]
        p0 = cell0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.line_spacing = 1.0
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9)
        r0.font.color.rgb = C_NAVY
        set_cell_background(cell0, "F1F5F9")
        set_cell_margins(cell0, top=40, bottom=40, left=80, right=80)

        cell1 = row.cells[1]
        cell1.width = col_widths[1]
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.line_spacing = 1.0
        r1 = p1.add_run(v)
        r1.font.size = Pt(9)
        r1.font.color.rgb = C_CHARCOAL
        set_cell_background(cell1, "F8FAFC")
        set_cell_margins(cell1, top=40, bottom=40, left=80, right=80)

    set_table_borders(table, color="CBD5E1", sz="4")
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(2.0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = C_NAVY
    return p

def add_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = C_TEAL
    return p

def _add_inline_runs(p, text):
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.font.bold = True
            run.font.color.rgb = C_CHARCOAL
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.font.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 83, 9)
        else:
            p.add_run(part)

def add_body(text, justified=True):
    p = doc.add_paragraph()
    if justified:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    _add_inline_runs(p, text)
    return p

def add_bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    _add_inline_runs(p, text)
    return p

def add_numbered(num_str, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    r_num = p.add_run(num_str + " ")
    r_num.font.bold = True
    r_num.font.color.rgb = C_NAVY
    _add_inline_runs(p, text)
    return p

def add_figure(img_path, caption_text, width_inches=5.4):
    if not os.path.exists(img_path):
        print(f"WARNING: image not found: {img_path}")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(3)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_inches))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(1)
    p_cap.paragraph_format.space_after = Pt(6)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = 'Calibri'
    run_cap.font.size = Pt(9)
    run_cap.font.bold = True
    run_cap.font.italic = True
    run_cap.font.color.rgb = C_MUTED

TABLE_COUNTER = {"n": 0}

def add_generic_table(headers, rows):
    TABLE_COUNTER["n"] += 1
    n_cols = len(headers)
    table = doc.add_table(rows=len(rows) + 1, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1.0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=55, bottom=55, left=70, right=70)

    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            if c_idx >= n_cols:
                break
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(2.5)
            p.paragraph_format.space_after = Pt(1.0)
            p.paragraph_format.line_spacing = 1.05
            parts = re.split(r'(\*\*.*?\*\*)', val)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                    run.font.size = Pt(8)
                    run.font.color.rgb = C_NAVY if c_idx == 0 else C_CHARCOAL
                else:
                    run = p.add_run(part)
                    run.font.size = Pt(8)
                    run.font.color.rgb = C_BODY
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=40, bottom=40, left=70, right=70)

    set_table_borders(table, color="CBD5E1", sz="4")
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(4)

with open('.docs/project-plan/final/project-plan-draft.md', 'r', encoding='utf-8') as f:
    lines = f.readlines()

add_title("AI MEETING COMMUNICATION COACH: A SPEAKER-AWARE DESKTOP ASSISTANT")
add_subtitle("Project Plan and Early-Stage Implementation Report")

metadata = {
    "Author / Candidate:": "Sneha Neupane (Student ID: 900008705)",
    "Unit of Study:": "ITNET312A Major Individual Project",
    "Assessment:": "Assessment 3: Project Plan / Early-Stage Presentation (Weighting: 15% Report + 5% Presentation)",
    "Institution:": "TAFE NSW Higher Education",
    "Submission Target:": "Week 8 via Moodle Turnitin",
    "Word Count:": "Approximately 2,000 words",
}
add_meta_box(metadata)

in_table = False
table_headers = []
table_rows = []

FIGURE_TRIGGER = "[Figure 1: End-to-end system architecture and dataflow block diagram]"

for raw_line in lines:
    line = raw_line.strip()

    if not line:
        continue

    is_table_line = line.startswith('|') and line.endswith('|')

    if is_table_line:
        cols = [c.strip() for c in line.split('|')[1:-1]]
        if all(re.match(r'^:?-+:?$', c) for c in cols):
            continue
        if not in_table:
            in_table = True
            table_headers = cols
            table_rows = []
        else:
            table_rows.append(cols)
        continue
    else:
        if in_table:
            in_table = False
            add_generic_table(table_headers, table_rows)
            table_headers = []
            table_rows = []

    if line.startswith('#### '):
        add_h2(line[5:])
        continue
    elif line.startswith('### '):
        add_h2(line[4:])
        continue
    elif line.startswith('## '):
        add_h1(line[3:])
        continue

    if FIGURE_TRIGGER in line:
        add_figure(
            ".docs/literature-review/final/figures/figure1_architecture.png",
            "Figure 1: End-to-End System Architecture and Dataflow Block Diagram.",
        )
        continue

    num_match = re.match(r'^(\d+\.)\s+(.*)', line)
    if num_match:
        add_numbered(num_match.group(1), num_match.group(2))
    elif line.startswith('* ') or line.startswith('- '):
        add_bullet(line[2:])
    else:
        add_body(line)

if in_table:
    add_generic_table(table_headers, table_rows)

output_path = ".docs/project-plan/final/project-plan.docx"
doc.save(output_path)
print(f"Successfully generated: {output_path}")
