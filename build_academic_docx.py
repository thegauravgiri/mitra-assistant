import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

doc = docx.Document()

# Standard 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    
    # Running Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("ITNET312A Major Individual Project | Literature Review")
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(9)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(80, 80, 80)
    
    # Running Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Sneha Neupane (900008705) • TAFE NSW Higher Education")
    frun.font.name = "Times New Roman"
    frun.font.size = Pt(9)
    frun.font.color.rgb = RGBColor(80, 80, 80)

# Strictly Academic Pure Black Palette
C_BLACK = RGBColor(0, 0, 0)
C_DARK_GRAY = RGBColor(60, 60, 60)

# Default style: Times New Roman, 12pt, Pure Black
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(12)
normal_style.font.color.rgb = C_BLACK
normal_style.paragraph_format.line_spacing = 1.15
normal_style.paragraph_format.space_after = Pt(5)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = C_BLACK
    return p

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = C_BLACK
    return p

def add_meta_box(meta_dict):
    table = doc.add_table(rows=len(meta_dict), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(2.0), Inches(4.5)]
    
    for i, (k, v) in enumerate(meta_dict.items()):
        row = table.rows[i]
        
        cell0 = row.cells[0]
        cell0.width = col_widths[0]
        p0 = cell0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.line_spacing = 1.05
        r0 = p0.add_run(k)
        r0.font.name = 'Times New Roman'
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = C_BLACK
        set_cell_background(cell0, "F3F4F6")
        set_cell_margins(cell0, top=45, bottom=45, left=80, right=80)
        
        cell1 = row.cells[1]
        cell1.width = col_widths[1]
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.line_spacing = 1.05
        r1 = p1.add_run(v)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_BLACK
        set_cell_background(cell1, "FFFFFF")
        set_cell_margins(cell1, top=45, bottom=45, left=80, right=80)

    set_table_borders(table, color="D1D5DB", sz="4")
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(10)

def add_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = C_BLACK
    return p

def add_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = C_BLACK
    return p

def add_h3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = C_BLACK
    return p

def add_body(text, justified=True):
    p = doc.add_paragraph()
    if justified:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(5)
    
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.color.rgb = C_BLACK
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Times New Roman'
            run.font.italic = True
            run.font.color.rgb = C_BLACK
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            run.font.color.rgb = C_BLACK
        else:
            run = p.add_run(part)
            run.font.name = 'Times New Roman'
            run.font.color.rgb = C_BLACK
    return p

def add_bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.color.rgb = C_BLACK
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Times New Roman'
            run.font.italic = True
            run.font.color.rgb = C_BLACK
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            run.font.color.rgb = C_BLACK
        else:
            run = p.add_run(part)
            run.font.name = 'Times New Roman'
            run.font.color.rgb = C_BLACK
    return p

def add_numbered(num_str, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    
    r_num = p.add_run(num_str + " ")
    r_num.font.name = 'Times New Roman'
    r_num.font.bold = True
    r_num.font.color.rgb = C_BLACK
    
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.color.rgb = C_BLACK
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Times New Roman'
            run.font.italic = True
            run.font.color.rgb = C_BLACK
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            run.font.color.rgb = C_BLACK
        else:
            run = p.add_run(part)
            run.font.name = 'Times New Roman'
            run.font.color.rgb = C_BLACK
    return p

def add_math_equation():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    
    omml_str = '''
    <m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
      <m:oMath>
        <m:r><m:rPr><m:scr m:val="roman"/><m:sty m:val="p"/></m:rPr><m:t>Participation Share</m:t></m:r>
        <m:r><m:rPr><m:scr m:val="roman"/><m:sty m:val="p"/></m:rPr><m:t> (%) = </m:t></m:r>
        <m:f>
          <m:fPr><m:type m:val="bar"/></m:fPr>
          <m:num>
            <m:r><m:rPr><m:scr m:val="roman"/><m:sty m:val="p"/></m:rPr><m:t>Total Speaking Time of Speaker i</m:t></m:r>
          </m:num>
          <m:den>
            <m:r><m:rPr><m:scr m:val="roman"/><m:sty m:val="p"/></m:rPr><m:t>Total Active Meeting Speaking Time</m:t></m:r>
          </m:den>
        </m:f>
        <m:r><m:rPr><m:scr m:val="roman"/><m:sty m:val="p"/></m:rPr><m:t> × 100</m:t></m:r>
      </m:oMath>
    </m:oMathPara>
    '''
    omml_element = parse_xml(omml_str)
    p._element.append(omml_element)

def add_figure(img_path, caption_text, width_inches=6.0):
    """Inserts a large, prominent, and properly readable figure."""
    if not os.path.exists(img_path):
        print(f"ERROR: Image {img_path} not found!")
        return
    print(f"Adding prominent figure: {img_path} (width={width_inches} in)")
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = 'Times New Roman'
    run_cap.font.size = Pt(10)
    run_cap.font.italic = True
    run_cap.font.color.rgb = C_BLACK

def add_reference(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    
    parts = re.split(r'(https?://[^\s]+)', text)
    for part in parts:
        if part.startswith('http'):
            run = p.add_run(part)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.color.rgb = C_BLACK
            run.font.underline = True
        else:
            subparts = re.split(r'(\*.*?\*)', part)
            for sp in subparts:
                if sp.startswith('*') and sp.endswith('*'):
                    r = p.add_run(sp[1:-1])
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(11)
                    r.font.italic = True
                    r.font.color.rgb = C_BLACK
                else:
                    r = p.add_run(sp)
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(11)
                    r.font.color.rgb = C_BLACK

def add_comparison_table(headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(1.3), Inches(1.3), Inches(1.3), Inches(1.3), Inches(1.3)]
    
    # Academic Header Row: Light gray fill, bold black text
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(title)
        run.font.name = 'Times New Roman'
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = C_BLACK
        set_cell_background(hdr_cells[i], "E5E7EB")
        set_cell_margins(hdr_cells[i], top=50, bottom=50, left=60, right=60)
        
    # Data rows: Clean alternating black text
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F9FAFB" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].width = col_widths[c_idx]
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            
            parts = re.split(r'(\*\*.*?\*\*)', val)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.name = 'Times New Roman'
                    run.font.bold = True
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = C_BLACK
                else:
                    run = p.add_run(part)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = C_BLACK
            
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=35, bottom=35, left=60, right=60)
            
    set_table_borders(table, color="9CA3AF", sz="4")
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(3)
    p_cap.paragraph_format.space_after = Pt(10)
    run_cap = p_cap.add_run("Table 1. State-of-the-Art Comparative Synthesis Matrix Across Meeting Assistant Paradigms.")
    run_cap.font.name = 'Times New Roman'
    run_cap.font.size = Pt(9.5)
    run_cap.font.italic = True
    run_cap.font.color.rgb = C_BLACK

# Load markdown draft
with open('.docs/literature-review/final/literature-review-draft.md', 'r', encoding='utf-8') as f:
    lines = f.readlines()

add_title("AI Meeting Communication Coach: A Speaker-Aware Desktop Assistant")
add_subtitle("Literature Review on Real-Time Conversational Co-Pilots, Streaming Diarisation, Sociometrics, and Cognitive Ergonomics")

metadata = {
    "Candidate:": "Sneha Neupane (Student ID: 900008705)",
    "Unit of Study:": "ITNET312A Major Individual Project",
    "Assessment:": "Assessment 2: Literature Review (Weighting: 20%)",
    "Institution:": "TAFE NSW Higher Education",
    "Submission Target:": "Week 6 via Moodle Turnitin",
    "Referencing Style:": "APA 7th Edition"
}
add_meta_box(metadata)

in_references = False
in_table = False
table_headers = []
table_rows = []

for line in lines:
    line = line.strip()
    
    if line.startswith('# AI Meeting Communication Coach') or line.startswith('**Author:**') or line.startswith('**Student ID:**') or line.startswith('**Course:**') or line.startswith('**Assessment:**') or line.startswith('**Institution:**') or line.startswith('**Date:**') or line == '---':
        continue
        
    if not line:
        continue
        
    if line.startswith('## 8. References'):
        if in_table:
            add_comparison_table(table_headers, table_rows)
            in_table = False
            table_headers = []
            table_rows = []
        in_references = True
        add_h1("8. References")
        continue
    elif line.startswith('## 9. Appendix'):
        in_references = False
        add_h1("9. Appendix: Generative AI Declaration")
        continue
    elif line.startswith('## '):
        in_references = False
        add_h1(line[3:])
        continue
    elif line.startswith('### '):
        in_references = False
        add_h2(line[4:])
        continue
    elif line.startswith('#### '):
        in_references = False
        add_h3(line[5:])
        continue
        
    # Table logic
    if line.startswith('|') and line.endswith('|'):
        cols = [c.strip() for c in line.split('|')[1:-1]]
        if all(re.match(r'^:?-+:?$', c) for c in cols):
            continue
        if not in_table:
            in_table = True
            table_headers = cols
            table_rows = []
        else:
            table_rows.append(cols)
        continue
    else:
        if in_table:
            in_table = False
            add_comparison_table(table_headers, table_rows)
            table_headers = []
            table_rows = []
            
    if in_references:
        add_reference(line)
        continue
        
    # Equation
    if "Participation Share (%) = (Total Speaking Duration of Speaker" in line:
        add_math_equation()
        continue
        
    num_match = re.match(r'^(\d+\.)\s+(.*)', line)
    if num_match:
        add_numbered(num_match.group(1), num_match.group(2))
    elif line.startswith('* ') or line.startswith('- '):
        add_bullet(line[2:])
    else:
        add_body(line)
    
    lower_line = line.lower()
    if "outlines the end-to-end dataflow pipeline of the proposed desktop assistant." in lower_line:
        add_figure(".docs/literature-review/final/figures/figure1_architecture.png", 
                   "Figure 1. End-to-end system pipeline and dataflow architecture of the proposed desktop assistant.",
                   width_inches=6.0)
    elif "therefore, assistive interfaces must remain compact, peripheral, and easily glanceable." in lower_line:
        add_figure(".docs/literature-review/final/figures/figure3_turn_taking.png", 
                   "Figure 2. Conversational dynamics, turn-taking transitions, and interruption detection model (adapting Sacks et al., 1974; Zimmerman & West, 1975; Stivers et al., 2009).",
                   width_inches=6.0)
    elif "packaged inside a native desktop overlay that requires no bot attendance." in lower_line:
        add_figure(".docs/literature-review/final/figures/figure2_comparison.png", 
                   "Figure 3. Architectural comparison: traditional cloud meeting bots vs. proposed desktop assistant overlay.",
                   width_inches=6.0)
    elif "global shortcuts without taking focus away from their meeting app." in lower_line:
        add_figure(".docs/literature-review/final/figures/figure4_glanceable_hud.png", 
                   "Figure 4. Glanceable HUD interaction model: compact ambient badge vs. expanded categorized insight cards.",
                   width_inches=6.0)

if in_table:
    add_comparison_table(table_headers, table_rows)

output_path = ".docs/literature-review/final/literature-review-sneha.docx"
doc.save(output_path)
print(f"Successfully generated academic docx: {output_path}")
